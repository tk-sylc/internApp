"""User-facing ledger workflow acceptance tests through authenticated HTTP APIs."""
from io import BytesIO
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.test import TestCase, Client
from openpyxl import load_workbook
from rest_framework.test import APIClient

from .models import ApprovedApplication


class LedgerAcceptanceTests(TestCase):
    def setUp(self):
        self.files = TemporaryDirectory()
        self.addCleanup(self.files.cleanup)
        settings = self.settings(APPROVED_LEDGER_OUTPUT_DIR=self.files.name)
        settings.enable()
        self.addCleanup(settings.disable)
        User = get_user_model()
        self.first = User.objects.create_user("first", email="first@example.co.jp", password="Test-Password-4839")
        self.second = User.objects.create_user("second", email="second@example.co.jp", password="Test-Password-5849")
        self.client = APIClient()
        self.client.force_authenticate(self.first)

    def create_entry(self):
        response = self.client.post("/api/approved-applications/", {
            "application_type": "pc", "operation_type": "purchase",
            "applicant_name": "確認用社員", "details": {"device_name": "確認用PC", "quantity": 1},
        }, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def test_other_operator_can_edit_with_immutable_creator_and_audit(self):
        record = self.create_entry()
        self.client.force_authenticate(self.second)
        response = self.client.patch(f"/api/approved-applications/{record['id']}/", {
            "revision": record["revision"], "applicant_name": "修正後の社員",
            "entered_by_name": "改ざん", "entered_by_email": "forged@example.co.jp",
        }, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        data = response.data
        self.assertEqual(data["entered_by_email"], self.first.email)
        self.assertEqual(data["details"], record["details"])
        detail = self.client.get(f"/api/approved-applications/{record['id']}/").data
        changes = [h for h in detail["history"] if h["action"] == "update"]
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0]["actor_email"], self.second.email)
        change = next(c for c in changes[0]["changes"] if c["field"] == "applicant_name")
        self.assertEqual(change["before"], "確認用社員")
        self.assertEqual(change["after"], "修正後の社員")

    def test_cancel_excludes_record_from_workbook_and_restore_keeps_history(self):
        record = self.create_entry()
        url = f"/api/approved-applications/{record['id']}/"
        self.client.force_authenticate(self.second)
        canceled = self.client.post(url + "cancel/", {"revision": record["revision"], "reason": "二重登録"}, format="json")
        self.assertEqual(canceled.status_code, 200, canceled.data)
        self.assertTrue(canceled.data["is_cancelled"])
        self.assertEqual(self.client.get("/api/ledgers/pc/").data["rows"], [])
        self.assertEqual(len(self.client.get("/api/ledgers/pc/?include_cancelled=1").data["rows"]), 1)
        download = self.client.get("/api/ledgers/pc/download/")
        self.assertEqual(download.status_code, 200)
        payload = b"".join(download.streaming_content)
        book = load_workbook(BytesIO(payload))
        self.assertEqual(book.active.max_row, 1)
        book.close()
        restored = self.client.post(url + "restore/", {"revision": canceled.data["revision"]}, format="json")
        self.assertEqual(restored.status_code, 200, restored.data)
        self.assertFalse(restored.data["is_cancelled"])
        self.assertEqual(ApprovedApplication.objects.count(), 1)
        history = self.client.get(url).data["history"]
        self.assertIn("cancel", [h["action"] for h in history])
        self.assertIn("restore", [h["action"] for h in history])

    def test_stale_operator_cannot_overwrite_or_cancel(self):
        record = self.create_entry()
        url = f"/api/approved-applications/{record['id']}/"
        updated = self.client.patch(url, {"revision": record["revision"], "notes": "新しいメモ"}, format="json")
        self.assertEqual(updated.status_code, 200, updated.data)
        self.client.force_authenticate(self.second)
        for suffix, method, data in [
            ("", self.client.patch, {"notes": "古い画面からのメモ"}),
            ("cancel/", self.client.post, {"reason": "古い画面からの取消"}),
        ]:
            response = method(url + suffix, {**data, "revision": record["revision"]}, format="json")
            self.assertEqual(response.status_code, 409, response.data)
        saved = ApprovedApplication.objects.get(pk=record["id"])
        self.assertEqual(saved.notes, "新しいメモ")
        self.assertFalse(saved.is_cancelled)

    def test_whole_ledger_has_all_columns_and_records_and_download_ignores_search(self):
        ApprovedApplication.objects.bulk_create([
            ApprovedApplication(application_type="pc", operation_type="purchase", applicant_name=f"社員{i}",
                entered_by=self.first, entered_by_name="登録担当者", entered_by_email=self.first.email,
                details={"device_name": f"PC-{i}", "quantity": i + 1}, notes="末尾の列も表示")
            for i in range(125)
        ])
        response = self.client.get("/api/ledgers/pc/")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(len(response.data["rows"]), 125)
        self.assertGreater(len(response.data["columns"]), 20)
        for row in response.data["rows"]:
            self.assertEqual(len(row["cells"]), len(response.data["columns"]))
        download = self.client.get("/api/ledgers/pc/download/?q=not-a-match&operation_type=return")
        self.assertEqual(download.status_code, 200)
        payload = b"".join(download.streaming_content)
        book = load_workbook(BytesIO(payload))
        self.assertEqual(book.active.max_row, 126)
        self.assertEqual([cell.value for cell in book.active[1]], [c["label"] for c in response.data["columns"]])
        book.close()

    def test_anonymous_cannot_read_history_download_or_retry(self):
        record = self.create_entry()
        client = APIClient()
        for path in ["/api/ledgers/", "/api/ledgers/pc/", "/api/ledgers/pc/download/", f"/api/approved-applications/{record['id']}/"]:
            self.assertIn(client.get(path).status_code, (401, 403), path)
        self.assertIn(client.post("/api/ledgers/pc/sync/", {}, format="json").status_code, (401, 403))

    def test_session_mutations_require_csrf(self):
        record = self.create_entry()
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.second)
        url = f"/api/approved-applications/{record['id']}/"
        self.assertEqual(client.patch(url, {"revision": record["revision"], "notes": "拒否"}, content_type="application/json").status_code, 403)
        self.assertEqual(client.post(url + "cancel/", {"revision": record["revision"], "reason": "拒否"}, content_type="application/json").status_code, 403)
        self.assertEqual(client.post("/api/ledgers/pc/sync/", {}, content_type="application/json").status_code, 403)

    def test_deleted_or_arbitrary_file_path_is_never_downloaded(self):
        self.assertEqual(self.client.get("/api/ledgers/not-a-ledger/download/").status_code, 404)
