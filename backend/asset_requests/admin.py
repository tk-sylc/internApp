import re
from datetime import datetime

from django.contrib import admin
from django.core.paginator import Paginator
from django.http import FileResponse, Http404
from django.urls import path, reverse
from django.template.response import TemplateResponse
from openpyxl import load_workbook

from accounts.models import UserProfile

from .approved_ledger_sync import FILE_NAMES, sync_approved_ledger

from .models import (
    ApprovedApplication,
    ExternalStorageRequest,
    LANRequest,
    PCRequest,
    SmartphoneRequest,
)


class BaseAssetRequestAdmin(admin.ModelAdmin):
    list_display = (
        "reference_number",
        "created_by",
        "requester_name",
        "department",
        "requester_email",
        "status",
        "created_at",
    )
    list_editable = ("status",)
    list_filter = ("status", "department", "created_at")
    search_fields = (
        "created_by__username",
        "requester_name",
        "department",
        "requester_email",
    )
    readonly_fields = ("reference_number", "created_by", "created_at", "updated_at")
    ordering = ("-created_at",)
    date_hierarchy = "created_at"
    list_select_related = ("created_by",)

    @admin.display(description="受付番号", ordering="id")
    def reference_number(self, obj):
        if obj is None:
            return "保存後に発行されます"
        return obj.reference_number

    def save_model(self, request, obj, form, change):
        if not change and not obj.created_by_id:
            obj.created_by = request.user
        super().save_model(request, obj, form, change)

    def get_search_results(self, request, queryset, search_term):
        """通常の検索に加えて、完全な受付番号でも申請を検索する。"""
        base_queryset = queryset
        queryset, use_distinct = super().get_search_results(
            request,
            queryset,
            search_term,
        )
        pattern = re.compile(
            rf"^{re.escape(self.model.REFERENCE_PREFIX)}-(\d{{8}})-(\d+)$",
            re.IGNORECASE,
        )
        match = pattern.fullmatch(search_term.strip())
        if not match:
            return queryset, use_distinct

        try:
            created_date = datetime.strptime(match.group(1), "%Y%m%d").date()
        except ValueError:
            return queryset, use_distinct

        reference_match = base_queryset.filter(
            pk=int(match.group(2)),
            created_at__date=created_date,
        )
        return queryset | reference_match, use_distinct


class PCRequestAdmin(BaseAssetRequestAdmin):
    list_display = BaseAssetRequestAdmin.list_display + (
        "management_number",
        "start_date",
    )
    search_fields = BaseAssetRequestAdmin.search_fields + (
        "applicant_name",
        "management_number",
    )
    list_filter = BaseAssetRequestAdmin.list_filter + ("start_date",)


class SmartphoneRequestAdmin(BaseAssetRequestAdmin):
    list_display = BaseAssetRequestAdmin.list_display + (
        "os",
        "model_name",
        "quantity",
        "purchase_date",
    )
    list_filter = BaseAssetRequestAdmin.list_filter + (
        "os",
        "line_type",
        "storage",
        "sim_required",
        "purchase_date",
    )
    search_fields = BaseAssetRequestAdmin.search_fields + ("model_name",)


class ExternalStorageRequestAdmin(BaseAssetRequestAdmin):
    list_display = BaseAssetRequestAdmin.list_display + (
        "device_name",
        "capacity",
        "loan_date",
    )
    search_fields = BaseAssetRequestAdmin.search_fields + (
        "applicant_name",
        "device_name",
    )
    list_filter = BaseAssetRequestAdmin.list_filter + ("loan_date",)


class LANRequestAdmin(BaseAssetRequestAdmin):
    list_display = BaseAssetRequestAdmin.list_display + (
        "device_type",
        "device_name",
        "quantity",
        "start_date",
        "return_date",
    )
    list_filter = BaseAssetRequestAdmin.list_filter + (
        "device_type",
        "start_date",
        "return_date",
    )
    search_fields = BaseAssetRequestAdmin.search_fields + ("device_name", "location")


@admin.register(ApprovedApplication)
class ApprovedApplicationAdmin(admin.ModelAdmin):
    change_list_template = (
        'admin/asset_requests/approvedapplication/change_list.html'
    )
    list_display = (
        "reference_number",
        "operation_type",
        "application_type",
        "applicant_name",
        "department",
        "entered_by_name",
        "entered_by_email",
        "created_at",
    )
    list_filter = (
        "operation_type",
        "application_type",
        "department",
        "created_at",
    )
    search_fields = (
        "applicant_name",
        "department",
        "entered_by_name",
        "entered_by_email",
        "entered_by__username",
    )
    readonly_fields = (
        "reference_number",
        "entered_by",
        "entered_by_name",
        "entered_by_email",
        "created_at",
        "updated_at",
    )
    exclude = ("approved_date",)
    date_hierarchy = "created_at"
    ordering = ("-created_at",)

    @admin.display(description="受付番号", ordering="id")
    def reference_number(self, obj):
        return obj.reference_number if obj else "保存後に発行されます"

    def get_urls(self):
        custom_urls = [
            path(
                'ledger/<str:application_type>/',
                self.admin_site.admin_view(self.preview_excel_ledger),
                name='asset_requests_approvedapplication_ledger',
            ),
            path(
                'excel/<str:application_type>/',
                self.admin_site.admin_view(self.download_excel_ledger),
                name='asset_requests_approvedapplication_excel',
            ),
        ]
        return custom_urls + super().get_urls()

    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context['ledger_preview_links'] = [
            {
                'label': self._ledger_label(value),
                'url': reverse(
                    'admin:asset_requests_approvedapplication_ledger',
                    args=(value,),
                ),
            }
            for value in FILE_NAMES
        ]
        return super().changelist_view(request, extra_context=extra_context)

    @staticmethod
    def _ledger_label(application_type):
        return (
            FILE_NAMES[application_type]
            .removeprefix('承認済み_')
            .removesuffix('管理台帳.xlsx')
        )

    def preview_excel_ledger(self, request, application_type):
        if not self.has_view_permission(request):
            raise Http404
        if application_type not in FILE_NAMES:
            raise Http404

        ledger_path = sync_approved_ledger(application_type)
        workbook = load_workbook(ledger_path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)
            headers = next(values, ())
            rows = list(values)
        finally:
            workbook.close()

        page_obj = Paginator(rows, 50).get_page(request.GET.get('page'))
        application_label = self._ledger_label(application_type)
        context = {
            **self.admin_site.each_context(request),
            'opts': self.model._meta,
            'title': f'{application_label} Excel台帳',
            'headers': headers,
            'page_obj': page_obj,
            'record_count': len(rows),
            'download_url': reverse(
                'admin:asset_requests_approvedapplication_excel',
                args=(application_type,),
            ),
        }
        return TemplateResponse(
            request,
            'admin/asset_requests/approvedapplication/ledger_preview.html',
            context,
        )

    def download_excel_ledger(self, request, application_type):
        if not self.has_view_permission(request):
            raise Http404
        if application_type not in FILE_NAMES:
            raise Http404

        ledger_path = sync_approved_ledger(application_type)
        return FileResponse(
            ledger_path.open('rb'),
            as_attachment=True,
            filename=FILE_NAMES[application_type],
        )

    def save_model(self, request, obj, form, change):
        if not change and not obj.entered_by_id:
            obj.entered_by = request.user
            try:
                obj.entered_by_name = request.user.profile.display_name.strip()
            except UserProfile.DoesNotExist:
                obj.entered_by_name = request.user.email or request.user.get_username()
            obj.entered_by_email = request.user.email
        super().save_model(request, obj, form, change)
