"""Executable API/DB/Excel checks corresponding to the user's debug workbook.
Uses the configured test database and temporary ledger files, never production data.
"""
from datetime import datetime
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from unittest import skipUnless
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient
from .models import ApprovedApplication, ApprovedApplicationHistory, ApprovedLedgerState
from .serializers import APPROVED_TYPE_DETAIL_FIELDS, APPROVED_OPERATION_DETAIL_FIELDS, APPROVED_COMMON_DETAIL_FIELDS

URL = "/api/approved-applications/"


class DebugChecklistTests(TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.setting = override_settings(APPROVED_LEDGER_OUTPUT_DIR=self.directory.name)
        self.setting.enable()
        self.addCleanup(self.setting.disable)
        self.user = get_user_model().objects.create_user("checklist", email="checklist@example.invalid", password="Check-Test-5839!")
        self.client = APIClient()
        self.client.force_login(self.user)

    def payload(self, kind="pc", operation="purchase"):
        keys = APPROVED_TYPE_DETAIL_FIELDS[kind] | APPROVED_OPERATION_DETAIL_FIELDS[operation] | APPROVED_COMMON_DETAIL_FIELDS
        if operation == "return":
            keys -= {"usage_start_date"}
        elif operation == "disposal":
            keys -= {"usage_start_date", "usage_end_date"}
        details = {key: "check-" + key for key in keys}
        details.update({key: "2026-09-10" for key in keys if key.endswith("_date")})
        details.update({key: "2" for key in keys if key in ("quantity", "cpu_ghz", "ram_gb")})
        if kind == "phone":
            details["phone_number"] = "09001234567"
        return {"application_type": kind, "operation_type": operation,
                "applicant_name": "debug-checklist", "department": "", "details": details,
                "notes": "first line\nsecond line"}

    def create(self, payload=None, key=None):
        options = {"HTTP_IDEMPOTENCY_KEY": str(key)} if key else {}
        response = self.client.post(URL, payload or self.payload(), format="json", **options)
        self.assertEqual(response.status_code, 201, response.data)
        return response.data

    def excel(self, kind):
        response = self.client.get(f"/api/ledgers/{kind}/download/")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(b"".join(response.streaming_content)), data_only=False)
        self.addCleanup(book.close)
        return book.active

    def check_combination(self, kind, operation):
        payload = self.payload(kind, operation)
        result = self.create(payload)
        saved = ApprovedApplication.objects.get(pk=result["id"])
        detail = self.client.get(f"{URL}{saved.pk}/").data
        self.assertEqual(saved.details, payload["details"])
        self.assertEqual(saved.notes, payload["notes"])
        self.assertEqual(detail["details"], saved.details)
        self.assertEqual(detail["reference_number"], saved.reference_number)
        self.assertEqual(detail["entered_by_email"], self.user.email)
        self.assertEqual(detail["created_at"], result["created_at"])
        self.assertEqual(saved.history.count(), 1)
        preview = self.client.get(f"/api/ledgers/{kind}/").data
        sheet = self.excel(kind)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual([c.value for c in sheet[1]], [c["label"] for c in preview["columns"]])
        displayed = [c.value.strftime("%Y-%m-%d %H:%M:%S") if isinstance(c.value, datetime) else c.value or "" for c in sheet[2]]
        self.assertEqual(displayed, [v if v is not None else "" for v in preview["rows"][0]["cells"]])
        self.assertEqual(preview["rows"][0]["cells"][0], saved.reference_number)

    def test_t001_lost_response_retry_returns_same_record_and_history(self):
        key = uuid4()
        original = self.create(key=key)  # server committed; discard this response
        state = ApprovedLedgerState.objects.get(pk="pc").generation
        retry = self.create(key=key)
        self.assertEqual(retry["id"], original["id"])
        self.assertEqual(retry["reference_number"], original["reference_number"])
        self.assertEqual(ApprovedApplication.objects.count(), 1)
        self.assertEqual(ApprovedApplicationHistory.objects.count(), 1)
        self.assertEqual(ApprovedLedgerState.objects.get(pk="pc").generation, state)
        self.assertEqual(self.excel("pc").max_row, 2)

    def test_t001_key_cannot_be_reused_for_different_payload_or_user(self):
        key = uuid4()
        self.create(key=key)
        changed = self.payload()
        changed["notes"] = "different"
        response = self.client.post(URL, changed, format="json", HTTP_IDEMPOTENCY_KEY=str(key))
        self.assertEqual(response.status_code, 409)
        other = get_user_model().objects.create_user("other", email="other@example.invalid")
        self.client.force_login(other)
        response = self.client.post(URL, self.payload(), format="json", HTTP_IDEMPOTENCY_KEY=str(key))
        self.assertEqual(response.status_code, 409)
        self.assertEqual(ApprovedApplication.objects.count(), 1)

    def test_t001_new_key_allows_intentional_identical_registration(self):
        self.create(key=uuid4())
        self.create(key=uuid4())
        self.assertEqual(ApprovedApplication.objects.count(), 2)

    def test_t001_invalid_key_rejected_without_saving(self):
        response = self.client.post(URL, self.payload(), format="json", HTTP_IDEMPOTENCY_KEY="invalid")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ApprovedApplication.objects.count(), 0)

    def test_t018_all_optional_fields_blank(self):
        result = self.create({"application_type": "pc", "operation_type": "purchase"})
        saved = ApprovedApplication.objects.get(pk=result["id"])
        self.assertEqual(saved.details, {})
        self.assertEqual(saved.applicant_name, "")
        self.assertEqual(saved.notes, "")
        self.assertEqual(self.excel("pc").max_row, 2)

    def test_t020_unicode_newlines_leading_zero_and_length_boundaries(self):
        payload = self.payload("phone")
        payload["applicant_name"] = "A" * 100
        payload["details"]["model_name"] = "device-\u65e5\u672c\u8a9e-\U0001f600\nline2"
        result = self.create(payload)
        saved = ApprovedApplication.objects.get(pk=result["id"])
        self.assertEqual(saved.details["phone_number"], "09001234567")
        values = [cell.value for cell in self.excel("phone")[2]]
        self.assertIn("09001234567", values)
        self.assertIn(payload["details"]["model_name"], values)
        self.assertIn(payload["notes"], values)
        payload["applicant_name"] = "A" * 101
        self.assertEqual(self.client.post(URL, payload, format="json").status_code, 400)
        payload["applicant_name"] = "within-limit"
        payload["details"]["model_name"] = "A" * 10000
        result = self.create(payload)
        self.assertIn("A" * 10000, [c.value for c in self.excel("phone")[2]] + [c.value for c in self.excel("phone")[3]])
        payload["details"]["model_name"] = "A" * 10001
        self.assertEqual(self.client.post(URL, payload, format="json").status_code, 400)
        response = self.client.patch(f"{URL}{result['id']}/", {"revision": 1, "details": payload["details"]}, format="json")
        self.assertEqual(response.status_code, 400)

    def test_t020_note_excel_cell_limit_prevents_silent_truncation(self):
        payload = self.payload()
        payload["notes"] = "A" * 32767
        record = self.create(payload)
        self.assertIn(payload["notes"], [c.value for c in self.excel("pc")[2]])
        payload["notes"] += "B"
        response = self.client.post(URL, payload, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("notes", response.data)
        changed = self.client.patch(f"{URL}{record['id']}/", {
            "revision": 1, "notes": payload["notes"]}, format="json")
        self.assertEqual(changed.status_code, 400)
        self.assertEqual(ApprovedApplication.objects.count(), 1)
        self.assertEqual(ApprovedApplication.objects.get().notes, "A" * 32767)
        self.assertEqual(ApprovedApplication.objects.get().revision, 1)

    def test_excel_timestamp_does_not_round_into_next_second(self):
        from datetime import timezone as dt_timezone
        record = self.create()
        stamp = datetime(2026, 9, 10, 14, 59, 59, 999999, tzinfo=dt_timezone.utc)
        ApprovedApplication.objects.filter(pk=record["id"]).update(created_at=stamp)
        preview = self.client.get("/api/ledgers/pc/").data
        index = next(i for i, col in enumerate(preview["columns"]) if col["key"] == "created_at")
        excel_value = self.excel("pc")[2][index].value
        self.assertEqual(excel_value.strftime("%Y-%m-%d %H:%M:%S"), preview["rows"][0]["cells"][index])
        self.assertEqual(excel_value.microsecond, 0)
        self.assertEqual(ApprovedApplication.objects.get().created_at, stamp)

    def test_t024_blank_reason_cancel_restore_preserves_all_history(self):
        record = self.create()
        path = f"{URL}{record['id']}/"
        for reason in ("", " ", "\n\t"):
            self.assertEqual(self.client.post(path+"cancel/", {"revision": 1, "reason": reason}, format="json").status_code, 400)
        canceled = self.client.post(path+"cancel/", {"revision": 1, "reason": "test"}, format="json")
        self.assertEqual(canceled.status_code, 200)
        self.assertEqual(len(self.client.get(URL).data), 0)
        self.assertEqual(len(self.client.get(URL+"?include_cancelled=1").data), 1)
        self.assertEqual(self.excel("pc").max_row, 1)
        restored = self.client.post(path+"restore/", {"revision": 2}, format="json")
        self.assertEqual(restored.status_code, 200)
        self.assertEqual(self.excel("pc").max_row, 2)
        self.assertEqual(ApprovedApplication.objects.get().history.count(), 3)
        self.assertEqual(restored.data["details"], record["details"])

    def test_t026_type_operation_change_and_explicit_clear(self):
        record = self.create(self.payload("pc", "loan"))
        path = f"{URL}{record['id']}/"
        changed = self.client.patch(path, {"revision": 1, "application_type": "phone", "operation_type": "disposal", "details": {"model_name": "phone", "disposal_date": "2026-09-10"}}, format="json")
        self.assertEqual(changed.status_code, 200)
        self.assertEqual(changed.data["details"]["device_name"], record["details"]["device_name"])
        self.assertEqual(self.excel("pc").max_row, 1)
        self.assertEqual(self.excel("phone").max_row, 2)
        cleared = self.client.patch(path, {"revision": 2, "details": {"device_name": ""}}, format="json")
        self.assertEqual(cleared.status_code, 200)
        self.assertNotIn("device_name", cleared.data["details"])

    def test_t027_write_failure_on_create_edit_cancel_then_retry(self):
        with patch("asset_requests.approved_ledger_sync._publish_workbook", side_effect=PermissionError("test output unavailable")):
            record = self.create()
            path = f"{URL}{record['id']}/"
            changed = self.client.patch(path, {"revision": 1, "notes": "updated"}, format="json")
            self.assertEqual(changed.status_code, 200)
            canceled = self.client.post(path+"cancel/", {"revision": 2, "reason": "test"}, format="json")
            self.assertEqual(canceled.status_code, 200)
            self.assertFalse(canceled.data["ledger_synced"])
        saved = ApprovedApplication.objects.get()
        self.assertEqual(saved.notes, "updated")
        self.assertTrue(saved.is_cancelled)
        self.assertEqual(saved.history.count(), 3)
        retry = self.client.post("/api/ledgers/pc/sync/", {}, format="json")
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(retry.data["sync"]["state"], "synced")
        self.assertEqual(self.excel("pc").max_row, 1)

    def test_t031_inactive_account_cannot_access_read_or_write(self):
        record = self.create()
        self.user.is_active = False
        self.user.save(update_fields=["is_active"])
        for url in (URL, f"{URL}{record['id']}/", "/api/ledgers/", "/api/ledgers/pc/", "/api/ledgers/pc/download/"):
            self.assertIn(self.client.get(url).status_code, (401, 403), url)
        for url, method, payload in (
            (URL, self.client.post, self.payload()),
            (f"{URL}{record['id']}/", self.client.patch, {"revision": 1, "notes": "unauthorized"}),
            (f"{URL}{record['id']}/cancel/", self.client.post, {"revision": 1, "reason": "test"}),
            (f"{URL}{record['id']}/restore/", self.client.post, {"revision": 1}),
            ("/api/ledgers/pc/sync/", self.client.post, {}),
        ):
            self.assertIn(method(url, payload, format="json").status_code, (401, 403), url)
        self.assertEqual(ApprovedApplication.objects.get().revision, 1)

    def test_t031_authentication_and_csrf_matrix(self):
        record = self.create()
        path = f"{URL}{record['id']}/"
        reads = (URL, path, "/api/ledgers/", "/api/ledgers/pc/", "/api/ledgers/pc/download/")
        writes = ((URL, "post", self.payload()),
                  (path, "patch", {"revision": 1, "notes": "unauthorized"}),
                  (path+"cancel/", "post", {"revision": 1, "reason": "test"}),
                  (path+"restore/", "post", {"revision": 1}),
                  ("/api/ledgers/pc/sync/", "post", {}))
        before = list(ApprovedLedgerState.objects.values())
        for mode in ("anonymous", "inactive", "missing_csrf", "invalid_csrf"):
            client = APIClient(enforce_csrf_checks=True)
            self.user.is_active = mode != "inactive"
            self.user.save(update_fields=["is_active"])
            if mode != "anonymous":
                client.force_login(self.user)
            if mode in ("anonymous", "inactive"):
                for url in reads:
                    with self.subTest(mode=mode, method="get", url=url):
                        self.assertIn(client.get(url).status_code, (401, 403))
            for url, method, data in writes:
                with self.subTest(mode=mode, method=method, url=url):
                    headers = {"HTTP_X_CSRFTOKEN": "invalid"} if mode == "invalid_csrf" else {}
                    self.assertIn(getattr(client, method)(url, data, format="json", **headers).status_code, (401, 403))
        self.assertEqual(ApprovedApplication.objects.count(), 1)
        self.assertEqual(ApprovedApplication.objects.get().revision, 1)
        self.assertEqual(ApprovedApplicationHistory.objects.count(), 1)
        self.assertEqual(list(ApprovedLedgerState.objects.values()), before)

    def test_t022_other_operator_changes_details_notes_and_excel_preserves_creator(self):
        original = self.create()
        saved = ApprovedApplication.objects.get()
        identity = (saved.entered_by_id, saved.entered_by_email, saved.entered_by_name, saved.created_at)
        other = get_user_model().objects.create_user("editor", email="editor@example.invalid")
        self.client.force_login(other)
        response = self.client.patch(f"{URL}{saved.pk}/", {
            "revision": 1, "details": {"device_name": "edited device"}, "notes": "edited note"}, format="json")
        self.assertEqual(response.status_code, 200)
        saved.refresh_from_db()
        self.assertEqual((saved.entered_by_id, saved.entered_by_email, saved.entered_by_name, saved.created_at), identity)
        values = [c.value for c in self.excel("pc")[2]]
        self.assertIn("edited device", values)
        self.assertIn("edited note", values)
        self.assertIn(self.user.email, values)
        detail = self.client.get(f"{URL}{saved.pk}/").data
        change = next(h for h in detail["history"] if h["action"] == "update")
        self.assertEqual(change["actor_email"], other.email)
        notes = next(item for item in change["changes"] if item["field"] == "notes")
        self.assertEqual(notes["before"], original["notes"])
        self.assertEqual(notes["after"], "edited note")

    @skipUnless(os.name == "nt", "uses real Windows workbook sharing locks")
    def test_t027_t028_real_workbook_lock_preserves_data_and_recovers_latest_download(self):
        import ctypes
        from ctypes import wintypes
        from .approved_ledger_sync import FILE_NAMES

        first = self.create()
        path = Path(self.directory.name) / FILE_NAMES["pc"]
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                     wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        kernel.CreateFileW.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel.CloseHandle.restype = wintypes.BOOL
        handle = kernel.CreateFileW(str(path), 0x80000000, 0, None, 3, 0, None)
        if handle == ctypes.c_void_p(-1).value:
            raise ctypes.WinError(ctypes.get_last_error())
        try:
            second = self.create()
            self.assertFalse(second["ledger_synced"])
            edited = self.client.patch(f"{URL}{first['id']}/", {
                "revision": 1, "notes": "latest after real lock"}, format="json")
            self.assertEqual(edited.status_code, 200)
            self.assertFalse(edited.data["ledger_synced"])
            canceled = self.client.post(f"{URL}{second['id']}/cancel/", {
                "revision": 1, "reason": "lock test"}, format="json")
            self.assertEqual(canceled.status_code, 200)
            self.assertFalse(canceled.data["ledger_synced"])
            failed = self.client.get("/api/ledgers/pc/download/")
            self.assertEqual(failed.status_code, 503)
            self.assertIn("application/json", failed["Content-Type"])
            self.assertEqual(ApprovedApplication.objects.count(), 2)
            self.assertEqual(ApprovedApplicationHistory.objects.count(), 4)
            self.assertEqual(ApprovedLedgerState.objects.get(pk="pc").state, "error")
        finally:
            kernel.CloseHandle(handle)
        retry = self.client.post("/api/ledgers/pc/sync/", {}, format="json")
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(retry.data["sync"]["state"], "synced")
        sheet = self.excel("pc")
        self.assertEqual(sheet.max_row, 2)
        values = [c.value for c in sheet[2]]
        self.assertIn("latest after real lock", values)
        self.assertIn(first["reference_number"], values)
        self.assertNotIn(second["reference_number"], values)


for index, (operation, kind) in enumerate(
    ((op, kind) for op in ("purchase", "loan", "return", "disposal") for kind in ("pc", "phone", "lan", "memory")), start=2
):
    def check(self, kind=kind, operation=operation):
        self.check_combination(kind, operation)
    check.__name__ = f"test_t{index:03d}_{operation}_{kind}"
    setattr(DebugChecklistTests, check.__name__, check)
