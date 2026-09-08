import json
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from threading import RLock
from contextlib import contextmanager
from datetime import datetime

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from accounts.models import UserProfile

from .models import ApprovedApplication, ApprovedLedgerState


DETAIL_COLUMNS = {
    "pc": (
        ("機種名", "device_name"),
        ("CPU（GHz）", "cpu_ghz"),
        ("RAM（GB）", "ram_gb"),
        ("OS・バージョン", "os"),
        ("セキュリティソフト", "security_software"),
        ("ウイルス対策ソフト導入確認", "antivirus_installed"),
        ("Officeバージョン", "office_version"),
        ("Browserバージョン", "browser_version"),
        ("Adobe Readerバージョン", "adobe_reader_version"),
        ("Flash Playerバージョン", "flash_player_version"),
    ),
    "memory": (
        ("外部記憶装置の種類", "storage_type"),
        ("機器名", "device_name"),
        ("容量", "capacity"),
        ("暗号化ソフト", "encryption_software"),
        ("ウイルスチェック", "virus_check"),
        ("ウイルスパターンファイル", "virus_pattern_file"),
    ),
    "lan": (
        ("機器種別", "device_type"),
        ("機器名", "device_name"),
        ("暗号方式", "wireless_encryption"),
        ("その他の暗号方式", "wireless_encryption_other"),
        ("入手方法", "acquisition_method"),
        ("借用元", "borrowed_from"),
    ),
    "phone": (
        ("OS・バージョン", "os"),
        ("機種", "model_name"),
        ("容量", "storage"),
        ("電話番号", "phone_number"),
        ("キャリア名", "carrier"),
        ("セキュリティソフト", "security_software"),
        ("ウイルス対策ソフト導入確認", "antivirus_installed"),
    ),
    "other": (("転記内容", "_all"),),
}

OPERATION_COLUMNS = (
    ("管理番号", "management_number"),
    ("利用開始日", "usage_start_date"),
    ("利用終了日", "usage_end_date"),
    ("利用場所", "location"),
    ("廃棄日", "disposal_date"),
    ("数量", "quantity"),
    ("目的", "purpose"),
    ("返却時の状態", "condition"),
    ("廃棄理由", "disposal_reason"),
    ("廃棄方法", "disposal_method"),
)

FILE_NAMES = {
    "pc": "承認済み_PC管理台帳.xlsx",
    "memory": "承認済み_外部記憶装置管理台帳.xlsx",
    "lan": "承認済み_LAN機器管理台帳.xlsx",
    "phone": "承認済み_スマートフォン管理台帳.xlsx",
    "other": "承認済み_その他申請管理台帳.xlsx",
}

# The process lock also serializes local SQLite development; in production the
# persistent row locks serialize all Gunicorn workers AND management commands.
_SYNC_LOCK = RLock()
LEDGER_LABELS = {"pc": "PC", "memory": "外部記憶装置", "lan": "LAN機器", "phone": "スマートフォン", "other": "その他"}
SYNC_ERROR_MESSAGE = "Excelへの反映に失敗しました。登録内容は保存されています。再試行してください。"


@contextmanager
def lock_ledgers(application_types):
    types = sorted(set(application_types))
    if any(value not in FILE_NAMES for value in types):
        raise ValueError("対応していない機器種別です。")
    with _SYNC_LOCK, transaction.atomic():
        # Rows are seeded by the migration. get_or_create also supports a fresh
        # test database after flush, without depending on process-local state.
        for value in types:
            ApprovedLedgerState.objects.get_or_create(application_type=value)
        states = list(ApprovedLedgerState.objects.select_for_update().filter(
            application_type__in=types,
        ).order_by("application_type"))
        yield states


def mark_pending(states):
    for state in states:
        state.generation += 1
        state.state = ApprovedLedgerState.State.PENDING
        state.error = ""
        state.save(update_fields=["generation", "state", "error"])


def sync_status(application_type):
    state, _ = ApprovedLedgerState.objects.get_or_create(application_type=application_type)
    return {
        "state": state.state,
        "synced_at": state.synced_at.isoformat() if state.synced_at else None,
        "error": state.error,
        "generation": state.generation,
        "synced_generation": state.synced_generation,
    }


def ledger_columns(application_type, records=()):
    columns = [
        {"key": "reference_number", "label": "受付番号"},
        {"key": "operation_type", "label": "処理区分"},
        {"key": "applicant_name", "label": "申請者氏名"},
        {"key": "department", "label": "所属部署"},
        {"key": "entered_by_name", "label": "登録担当者"},
        {"key": "created_at", "label": "登録日時"},
    ]
    detail_columns = (*DETAIL_COLUMNS[application_type], *OPERATION_COLUMNS)
    columns.extend({"key": "details." + key, "label": label} for label, key in detail_columns)
    known = {key for _label, key in detail_columns}
    # Older records may contain fields from another type or an older form.
    # Include them in BOTH the full preview and export; never hide stored data.
    labels = {key: label for cols in DETAIL_COLUMNS.values() for label, key in cols}
    labels.update({key: label for label, key in OPERATION_COLUMNS})
    if "_all" not in known:
        extras = sorted({key for record in records for key in record.details} - known)
        columns.extend({"key": "details." + key, "label": _safe_value(labels.get(key, key))} for key in extras)
    columns.append({"key": "notes", "label": "担当者メモ"})
    return columns


def record_cells(record, columns):
    values = []
    for column in columns:
        key = column["key"]
        if key.startswith("details."):
            field = key.removeprefix("details.")
            value = record.details if field == "_all" else record.details.get(field, "")
        elif key == "operation_type":
            value = record.get_operation_type_display()
        elif key == "entered_by_name":
            value = _operator_name(record)
        elif key == "created_at":
            value = _excel_datetime(record.created_at)
        else:
            value = getattr(record, key)
        values.append(_safe_value(value))
    return values


def ledger_records(application_type, include_cancelled=False):
    queryset = ApprovedApplication.objects.filter(application_type=application_type)
    if not include_cancelled:
        queryset = queryset.filter(is_cancelled=False)
    return list(queryset.select_related("entered_by__profile").order_by("pk"))


def ledger_preview(application_type, include_cancelled=False):
    # A coherent snapshot of rows, columns, and state even during a concurrent
    # edit; mutations use this same lock order.
    with lock_ledgers([application_type]):
        records = ledger_records(application_type, include_cancelled)
        columns = ledger_columns(application_type, records)
        return {
            "application_type": application_type,
            "label": LEDGER_LABELS[application_type],
            "columns": columns,
            "rows": [
                {"id": record.pk, "is_cancelled": record.is_cancelled,
                 "cells": [value.strftime("%Y-%m-%d %H:%M:%S") if isinstance(value, datetime)
                           else str(value) if value is not None else ""
                           for value in record_cells(record, columns)]}
                for record in records
            ],
            "sync": sync_status(application_type),
        }



def _safe_value(value):
    if isinstance(value, bool):
        return "あり" if value else "なし"
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _excel_datetime(value):
    """Excelが扱えないタイムゾーン情報を外し、設定中の現地時刻にする。"""

    if timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def _operator_name(record):
    if record.entered_by_name:
        return record.entered_by_name
    try:
        display_name = record.entered_by.profile.display_name.strip()
    except UserProfile.DoesNotExist:
        display_name = ""
    return display_name or record.entered_by.email or record.entered_by.get_username()


def _publish_workbook(application_type):
    """Called only with the ledger row lock held for snapshot through replace."""
    output_dir = Path(settings.APPROVED_LEDGER_OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / FILE_NAMES[application_type]
    records = ledger_records(application_type)
    columns = ledger_columns(application_type, records)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "転記データ"
    worksheet.freeze_panes = "A2"
    worksheet.append([column["label"] for column in columns])
    header_fill = PatternFill(fill_type="solid", fgColor="17376D")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for record in records:
        worksheet.append(record_cells(record, columns))
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        longest = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(longest + 3, 42)
    temporary_path = None
    try:
        with NamedTemporaryFile(dir=output_dir, prefix=f".{output_path.stem}-", suffix=".tmp.xlsx", delete=False) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if temporary_path and temporary_path.exists():
            temporary_path.unlink()
    return output_path


def sync_approved_ledger(application_type):
    failure = None
    with lock_ledgers([application_type]) as states:
        state = states[0]
        try:
            output_path = _publish_workbook(application_type)
        except Exception as exc:
            # Persist failure outside the failed generation operation; expose
            # no filesystem path, DB credentials, or underlying exception.
            state.state = ApprovedLedgerState.State.ERROR
            state.error = SYNC_ERROR_MESSAGE
            state.save(update_fields=["state", "error"])
            failure = exc
        else:
            state.state = ApprovedLedgerState.State.SYNCED
            state.synced_generation = state.generation
            state.synced_at = timezone.now()
            state.error = ""
            state.save(update_fields=["state", "synced_generation", "synced_at", "error"])
    if failure is not None:
        raise failure
    return output_path
