from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from accounts.models import UserProfile, Department
from .approved_ledger_sync import FILE_NAMES, sync_approved_ledger
from .models import ApprovedApplication, ApprovedApplicationHistory, ApprovedLedgerState


class LedgerWorkflowTests(APITestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.output = Path(self.directory.name)
        setting = override_settings(APPROVED_LEDGER_OUTPUT_DIR=self.output)
        setting.enable()
        self.addCleanup(setting.disable)
        self.user = get_user_model().objects.create_user(username="operator", email="operator@example.com", password="Strong-test-123!")
        UserProfile.objects.create(user=self.user, display_name="登録担当者", department=Department.SYSTEM)
        self.client.force_login(self.user)
        self.list_url = reverse("asset_requests:approved-application-list-create")

    def create_entry(self, **overrides):
        payload = {"application_type": "pc", "operation_type": "loan", "applicant_name": "申請者", "details": {"device_name": "テストPC", "quantity": 1}}
        payload.update(overrides)
        response = self.client.post(self.list_url, payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        return response

    def detail_url(self, pk):
        return reverse("asset_requests:approved-application-detail", args=[pk])

    def ledger_url(self, name, kind="pc"):
        return reverse("asset_requests:ledger-" + name, args=[kind])

    def test_sync_failure_preserves_registered_data_and_retry_recovers(self):
        with patch("asset_requests.approved_ledger_sync._publish_workbook", side_effect=PermissionError("PRIVATE_PATH")):
            response = self.create_entry()
        self.assertFalse(response.data["ledger_synced"])
        self.assertNotIn("PRIVATE_PATH", str(response.data))
        self.assertEqual(ApprovedApplication.objects.count(), 1)
        self.assertEqual(ApprovedApplicationHistory.objects.count(), 1)
        self.assertEqual(ApprovedLedgerState.objects.get(pk="pc").state, "error")
        retry = self.client.post(self.ledger_url("sync"))
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(retry.data["sync"]["state"], "synced")
        self.assertEqual(ApprovedApplication.objects.count(), 1)
        state = ApprovedLedgerState.objects.get(pk="pc")
        self.assertEqual(state.generation, state.synced_generation)
        self.assertEqual(state.error, "")

    def test_patch_preserves_omitted_and_legacy_details_and_creator(self):
        response = self.create_entry()
        record = ApprovedApplication.objects.get()
        record.details["old_custom_field"] = "過去のデータ"
        record.save(update_fields=["details"])
        updated = self.client.patch(self.detail_url(record.pk), {"revision": 1, "notes": "メモのみ修正", "entered_by_name": "改ざん", "is_cancelled": True}, format="json")
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.data["details"], record.details)
        self.assertEqual(updated.data["entered_by_name"], "登録担当者")
        self.assertFalse(updated.data["is_cancelled"])
        self.assertEqual(updated.data["revision"], 2)
        changes = updated.data["history"][0]["changes"]
        self.assertEqual([item["field"] for item in changes], ["notes"])

    def test_type_change_rebuilds_both_ledgers_and_preserves_historical_fields(self):
        entry = self.create_entry()
        updated = self.client.patch(self.detail_url(entry.data["id"]), {"revision": 1, "application_type": "phone", "details": {"model_name": "電話機"}}, format="json")
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.data["details"]["device_name"], "テストPC")
        old = load_workbook(self.output / FILE_NAMES["pc"])
        self.addCleanup(old.close)
        self.assertEqual(old.active.max_row, 1)
        new = load_workbook(self.output / FILE_NAMES["phone"])
        self.addCleanup(new.close)
        self.assertEqual(new.active.max_row, 2)
        values = [cell.value for cell in new.active[2]]
        self.assertIn("電話機", values)
        self.assertIn("テストPC", values)
        for value in ("pc", "phone"):
            self.assertEqual(ApprovedLedgerState.objects.get(pk=value).state, "synced")

    def test_malformed_patch_cannot_erase_data(self):
        entry = self.create_entry()
        for invalid in ([], "incorrect", None, {"device_name": {"injected": "nested"}}):
            with self.subTest(invalid=invalid):
                response = self.client.patch(self.detail_url(entry.data["id"]), {"revision": 1, "details": invalid}, format="json")
                self.assertEqual(response.status_code, 400)
        record = ApprovedApplication.objects.get()
        self.assertEqual(record.revision, 1)
        self.assertEqual(record.details["device_name"], "テストPC")
        self.assertEqual(record.history.count(), 1)

    def test_audit_failure_rolls_back_record_revision_and_dirty_state(self):
        entry = self.create_entry()
        state_before = ApprovedLedgerState.objects.get(pk="pc").generation
        with patch("asset_requests.approved_workflow._history", side_effect=RuntimeError("audit unavailable")):
            with self.assertRaises(RuntimeError):
                self.client.patch(self.detail_url(entry.data["id"]), {"revision": 1, "notes": "変更しない"}, format="json")
        record = ApprovedApplication.objects.get()
        self.assertEqual(record.notes, "")
        self.assertEqual(record.revision, 1)
        self.assertEqual(record.history.count(), 1)
        self.assertEqual(ApprovedLedgerState.objects.get(pk="pc").generation, state_before)

    def test_download_failure_never_serves_previous_file(self):
        self.create_entry()
        stale = (self.output / FILE_NAMES["pc"]).read_bytes()
        with patch("asset_requests.approved_ledger_sync._publish_workbook", side_effect=OSError("failed")):
            response = self.client.get(self.ledger_url("download"))
        self.assertEqual(response.status_code, 503)
        self.assertNotIn("attachment", response.headers.get("Content-Disposition", ""))
        self.assertEqual((self.output / FILE_NAMES["pc"]).read_bytes(), stale)
        self.assertEqual(ApprovedLedgerState.objects.get(pk="pc").state, "error")

    def test_download_read_failure_marks_error(self):
        self.create_entry()
        with patch("pathlib.Path.read_bytes", side_effect=OSError("failed")):
            response = self.client.get(self.ledger_url("download"))
        self.assertEqual(response.status_code, 503)
        self.assertEqual(ApprovedLedgerState.objects.get(pk="pc").state, "error")

    def test_atomic_publish_keeps_previous_file_if_save_fails(self):
        self.create_entry()
        previous = (self.output / FILE_NAMES["pc"]).read_bytes()
        with patch("asset_requests.approved_ledger_sync.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                sync_approved_ledger("pc")
        self.assertEqual((self.output / FILE_NAMES["pc"]).read_bytes(), previous)
        self.assertFalse(list(self.output.glob("*.tmp.xlsx")))
        self.assertEqual(ApprovedLedgerState.objects.get(pk="pc").state, "error")

    def test_formula_payload_remains_text_and_preview_matches_headers(self):
        self.create_entry(applicant_name="=HYPERLINK(\"https://example.invalid\")", details={"device_name": "+1+1", "quantity": 2})
        preview = self.client.get(self.ledger_url("detail"))
        response = self.client.get(self.ledger_url("download"))
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
        self.addCleanup(workbook.close)
        self.assertEqual([cell.value for cell in workbook.active[1]], [column["label"] for column in preview.data["columns"]])
        cell = workbook.active.cell(row=2, column=3)
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("'="))
        self.assertEqual(preview.data["rows"][0]["cells"][2], cell.value)

    def test_admin_business_mutation_is_disabled(self):
        entry = self.create_entry()
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        change = reverse("admin:asset_requests_approvedapplication_change", args=[entry.data["id"]])
        self.assertEqual(self.client.get(change).status_code, 200)
        self.assertEqual(self.client.post(change, {"notes": "bypass"}).status_code, 403)
        self.assertEqual(self.client.post(reverse("admin:asset_requests_approvedapplication_add"), {}).status_code, 403)
        self.assertEqual(self.client.post(reverse("admin:asset_requests_approvedapplication_delete", args=[entry.data["id"]]), {"post": "yes"}).status_code, 403)
        self.assertEqual(ApprovedApplication.objects.get().revision, 1)

    def test_sync_status_and_history_responses_are_not_cached(self):
        entry = self.create_entry()
        for url in (self.detail_url(entry.data["id"]), self.ledger_url("detail"), reverse("asset_requests:ledger-list")):
            response = self.client.get(url)
            self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.client.logout()
        response = self.client.get(self.detail_url(entry.data["id"]))
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.data["code"], "authentication_required")

    def test_dynamic_legacy_column_header_cannot_be_an_excel_formula(self):
        entry = self.create_entry()
        updated = self.client.patch(self.detail_url(entry.data["id"]), {"revision": 1, "details": {"=1+1": "古い追加項目"}}, format="json")
        self.assertEqual(updated.status_code, 200)
        preview = self.client.get(self.ledger_url("detail"))
        workbook = load_workbook(self.output / FILE_NAMES["pc"])
        self.addCleanup(workbook.close)
        header = next(cell for cell in workbook.active[1] if cell.value == "'=1+1")
        self.assertEqual(header.data_type, "s")
        self.assertIn("'=1+1", [column["label"] for column in preview.data["columns"]])

    def test_unchanged_legacy_department_can_be_saved_but_new_invalid_value_rejected(self):
        entry = self.create_entry()
        ApprovedApplication.objects.filter(pk=entry.data["id"]).update(department="旧営業本部")
        response = self.client.patch(self.detail_url(entry.data["id"]), {"revision": 1, "department": "旧営業本部", "notes": "メモを更新"}, format="json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["department"], "旧営業本部")
        response = self.client.patch(self.detail_url(entry.data["id"]), {"revision": 2, "department": "存在しない新部署"}, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ApprovedApplication.objects.get().revision, 2)
