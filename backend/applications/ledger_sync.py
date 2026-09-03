import os
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from threading import Lock

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import (
    ExternalStorageLoanApplication,
    LanEquipmentLoanApplication,
    PcLoanApplication,
    SmartphonePurchaseApplication,
)


COMMON_COLUMNS = (
    ('申請者氏名', 'requester_name'),
    ('所属部署', 'department'),
    ('社員番号', 'employee_number'),
)

LEDGER_CONFIG = {
    'pc': {
        'model': PcLoanApplication,
        'columns': COMMON_COLUMNS + (
            ('貸出者氏名', 'applicant_name'),
            ('管理番号', 'management_number'),
            ('利用開始日', 'start_date'),
            ('利用場所', 'location'),
        ),
        'sheet_name': 'PC貸出',
        'file_name': 'PC貸出管理台帳.xlsx',
    },
    'memory': {
        'model': ExternalStorageLoanApplication,
        'columns': COMMON_COLUMNS + (
            ('貸出者氏名', 'applicant_name'),
            ('機器名', 'device_name'),
            ('容量', 'capacity'),
            ('場所', 'location'),
            ('貸し出し日', 'loan_date'),
        ),
        'sheet_name': '外部記憶装置貸出',
        'file_name': '外部記憶装置貸出管理台帳.xlsx',
    },
    'lan': {
        'model': LanEquipmentLoanApplication,
        'columns': COMMON_COLUMNS + (
            ('機器種別', 'device_type'),
            ('機器名', 'device_name'),
            ('利用開始日', 'start_date'),
            ('返却予定日', 'return_date'),
            ('利用場所', 'location'),
        ),
        'sheet_name': 'LAN機器貸出',
        'file_name': 'LAN機器貸出管理台帳.xlsx',
    },
    'phone': {
        'model': SmartphonePurchaseApplication,
        'columns': COMMON_COLUMNS + (
            ('機種', 'model_name'),
            ('購入日', 'purchase_date'),
            ('容量', 'storage'),
            ('SIMの有無', 'sim_required'),
        ),
        'sheet_name': 'スマートフォン購入',
        'file_name': 'スマートフォン購入管理台帳.xlsx',
    },
}

_SYNC_LOCK = Lock()


def _safe_excel_value(value):
    """入力文字列がExcelの数式として実行されることを防ぐ。"""

    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return f"'{value}"
    return value


def _create_workbook(queryset, columns, sheet_name):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.freeze_panes = 'A2'

    worksheet.append([label for label, _field_name in columns])
    header_fill = PatternFill(fill_type='solid', fgColor='17376D')
    for cell in worksheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for application in queryset:
        values = [
            _safe_excel_value(getattr(application, field_name))
            for _label, field_name in columns
        ]
        worksheet.append(values)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = 'yyyy-mm-dd'

    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        maximum_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            maximum_length + 4,
            40,
        )

    return workbook


def sync_ledger(request_type):
    """SQLiteの全申請から指定された種類のExcel台帳を安全に再生成する。"""

    config = LEDGER_CONFIG.get(request_type)
    if config is None:
        raise ValueError(f'未対応の申請種別です: {request_type}')

    output_dir = Path(settings.LEDGER_OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / config['file_name']

    with _SYNC_LOCK:
        queryset = config['model'].objects.all().order_by('pk')
        workbook = _create_workbook(
            queryset,
            config['columns'],
            config['sheet_name'],
        )

        temporary_path = None
        try:
            with NamedTemporaryFile(
                dir=output_dir,
                prefix=f'.{output_path.stem}-',
                suffix='.tmp.xlsx',
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)

            workbook.save(temporary_path)
            os.replace(temporary_path, output_path)
        finally:
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()

    return output_path


def sync_all_ledgers():
    """4種類すべてのExcel台帳を再生成する。"""

    return {
        request_type: sync_ledger(request_type)
        for request_type in LEDGER_CONFIG
    }
