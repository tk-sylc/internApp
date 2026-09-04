import json
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from .ledger_sync import sync_all_ledgers

from .models import (
    ApprovedApplication,
    ApplicationStatus,
    ExternalStorageLoanApplication,
    LanEquipmentLoanApplication,
    PcLoanApplication,
    SmartphonePurchaseApplication,
)


User = get_user_model()


COMMON_FIELDS = {
    'requester_name': '山田 太郎',
    'department': '営業部',
    'employee_number': 'EMP-0124',
}


class ApplicationModelTests(TestCase):
    def test_save_all_application_types(self):
        PcLoanApplication.objects.create(
            **COMMON_FIELDS,
            applicant_name='山田 太郎',
            management_number='PC-01234',
            start_date=date(2026, 9, 3),
            location='東京本社',
        )
        ExternalStorageLoanApplication.objects.create(
            **COMMON_FIELDS,
            applicant_name='山田 太郎',
            device_name='USBメモリ',
            capacity='64GB',
            location='東京本社',
            loan_date=date(2026, 9, 3),
        )
        LanEquipmentLoanApplication.objects.create(
            **COMMON_FIELDS,
            device_type=LanEquipmentLoanApplication.DeviceType.LAN_CABLE,
            device_name='CAT6 LANケーブル',
            start_date=date(2026, 9, 3),
            return_date=date(2026, 9, 10),
            location='第2会議室',
        )
        SmartphonePurchaseApplication.objects.create(
            **COMMON_FIELDS,
            model_name='iPhone 16',
            purchase_date=date(2026, 9, 3),
            storage=SmartphonePurchaseApplication.Storage.GB_128,
            sim_required=SmartphonePurchaseApplication.SimRequired.YES,
        )

        self.assertEqual(PcLoanApplication.objects.count(), 1)
        self.assertEqual(ExternalStorageLoanApplication.objects.count(), 1)
        self.assertEqual(LanEquipmentLoanApplication.objects.count(), 1)
        self.assertEqual(SmartphonePurchaseApplication.objects.count(), 1)
        self.assertEqual(
            PcLoanApplication.objects.get().status,
            ApplicationStatus.PENDING,
        )

    def test_lan_return_date_must_not_precede_start_date(self):
        application = LanEquipmentLoanApplication(
            **COMMON_FIELDS,
            device_type=LanEquipmentLoanApplication.DeviceType.LAN_CABLE,
            device_name='CAT6 LANケーブル',
            start_date=date(2026, 9, 10),
            return_date=date(2026, 9, 3),
            location='第2会議室',
        )

        with self.assertRaises(ValidationError):
            application.full_clean()

    def test_application_status_can_be_updated(self):
        application = PcLoanApplication.objects.create(
            **COMMON_FIELDS,
            applicant_name='山田 太郎',
            management_number='PC-01234',
            start_date=date(2026, 9, 3),
            location='東京本社',
        )

        application.status = ApplicationStatus.APPROVED
        application.save(update_fields=['status'])
        application.refresh_from_db()

        self.assertEqual(application.status, ApplicationStatus.APPROVED)


class CreateApplicationViewTests(TestCase):
    def setUp(self):
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.ledger_directory = Path(self.temporary_directory.name)
        self.ledger_settings = override_settings(
            LEDGER_OUTPUT_DIR=self.ledger_directory,
        )
        self.ledger_settings.enable()
        self.addCleanup(self.ledger_settings.disable)

        self.client = Client(enforce_csrf_checks=True)
        self.user = User.objects.create_user(username='asset-manager', password='test-password')
        self.client.force_login(self.user)
        self.csrf_url = reverse('applications:csrf-token')
        self.create_url = reverse('applications:create-application')
        self.payload = {
            'requestType': 'pc',
            'requesterName': '山田 太郎',
            'department': '営業部',
            'employeeNumber': 'EMP-0124',
            'applicantName': '山田 太郎',
            'managementNumber': 'PC-01234',
            'startDate': '2026-09-03',
            'location': '東京本社',
        }

    def post_with_csrf(self, payload):
        self.client.get(self.csrf_url)
        token = self.client.cookies['csrftoken'].value
        return self.client.post(
            self.create_url,
            data=json.dumps(payload),
            content_type='application/json',
            HTTP_X_CSRFTOKEN=token,
        )

    def test_create_pc_application(self):
        response = self.post_with_csrf(self.payload)

        self.assertEqual(response.status_code, 201)
        self.assertEqual(PcLoanApplication.objects.count(), 1)
        application = PcLoanApplication.objects.get()
        self.assertEqual(application.management_number, 'PC-01234')
        self.assertEqual(response.json()['id'], application.pk)
        self.assertEqual(response.json()['status'], ApplicationStatus.PENDING)
        self.assertTrue(response.json()['ledgerSynced'])

        workbook = load_workbook(self.ledger_directory / 'PC貸出管理台帳.xlsx')
        worksheet = workbook['PC貸出']
        headers = tuple(cell.value for cell in worksheet[1])
        self.assertEqual(
            headers,
            (
                '申請者氏名',
                '所属部署',
                '社員番号',
                '貸出者氏名',
                '管理番号',
                '利用開始日',
                '利用場所',
            ),
        )

    def test_create_each_non_pc_application_type(self):
        cases = [
            (
                {
                    'requestType': 'memory',
                    'requesterName': '山田 太郎',
                    'department': '営業部',
                    'employeeNumber': 'EMP-0124',
                    'applicantName': '山田 太郎',
                    'deviceName': 'USBメモリ',
                    'capacity': '64GB',
                    'location': '東京本社',
                    'loanDate': '2026-09-03',
                },
                ExternalStorageLoanApplication,
            ),
            (
                {
                    'requestType': 'lan',
                    'requesterName': '山田 太郎',
                    'department': '営業部',
                    'employeeNumber': 'EMP-0124',
                    'deviceType': 'LANケーブル',
                    'deviceName': 'CAT6 LANケーブル',
                    'startDate': '2026-09-03',
                    'returnDate': '2026-09-10',
                    'location': '第2会議室',
                },
                LanEquipmentLoanApplication,
            ),
            (
                {
                    'requestType': 'phone',
                    'requesterName': '山田 太郎',
                    'department': '営業部',
                    'employeeNumber': 'EMP-0124',
                    'model': 'iPhone 16',
                    'deliveryDate': '2026-09-03',
                    'storage': '128GB',
                    'simRequired': 'あり',
                },
                SmartphonePurchaseApplication,
            ),
        ]

        for payload, model in cases:
            with self.subTest(request_type=payload['requestType']):
                response = self.post_with_csrf(payload)
                self.assertEqual(response.status_code, 201)
                self.assertEqual(model.objects.count(), 1)

    def test_reject_missing_required_field(self):
        del self.payload['managementNumber']

        response = self.post_with_csrf(self.payload)

        self.assertEqual(response.status_code, 400)
        self.assertEqual(PcLoanApplication.objects.count(), 0)

    def test_application_remains_saved_when_ledger_sync_fails(self):
        blocked_output_path = self.ledger_directory / 'not-a-directory'
        blocked_output_path.write_text('block directory creation', encoding='utf-8')

        with self.assertLogs('applications.views', level='ERROR'):
            with override_settings(LEDGER_OUTPUT_DIR=blocked_output_path):
                response = self.post_with_csrf(self.payload)

        self.assertEqual(response.status_code, 201)
        self.assertEqual(PcLoanApplication.objects.count(), 1)
        self.assertFalse(response.json()['ledgerSynced'])
        self.assertIsNotNone(response.json()['ledgerWarning'])

    def test_reject_post_without_csrf_token(self):
        response = self.client.post(
            self.create_url,
            data=json.dumps(self.payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 403)
        self.assertEqual(PcLoanApplication.objects.count(), 0)


class AuthenticationViewTests(TestCase):
    def setUp(self):
        self.client = Client(enforce_csrf_checks=True)
        self.user = User.objects.create_user(
            username='asset-manager',
            email='manager@example.com',
            password='test-password',
            first_name='資産管理',
            last_name='担当者',
        )

    def test_session_login_and_logout(self):
        session_url = reverse('applications:session')
        response = self.client.get(session_url)
        self.assertFalse(response.json()['authenticated'])
        token = self.client.cookies['csrftoken'].value

        response = self.client.post(
            reverse('applications:login'),
            data=json.dumps({'username': 'manager@example.com', 'password': 'test-password'}),
            content_type='application/json',
            HTTP_X_CSRFTOKEN=token,
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['authenticated'])
        self.assertEqual(response.json()['user']['username'], 'asset-manager')

        token = self.client.cookies['csrftoken'].value
        response = self.client.post(
            reverse('applications:logout'),
            HTTP_X_CSRFTOKEN=token,
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()['authenticated'])

    def test_approved_applications_require_login(self):
        response = self.client.get(reverse('applications:approved-applications'))
        self.assertEqual(response.status_code, 401)


class ApprovedApplicationViewTests(TestCase):
    def setUp(self):
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.ledger_directory = Path(self.temporary_directory.name)
        self.ledger_settings = override_settings(
            APPROVED_LEDGER_OUTPUT_DIR=self.ledger_directory,
        )
        self.ledger_settings.enable()
        self.addCleanup(self.ledger_settings.disable)
        self.user = User.objects.create_user(username='asset-manager', password='test-password')
        self.client = Client(enforce_csrf_checks=True)
        self.client.force_login(self.user)
        self.url = reverse('applications:approved-applications')
        self.payload = {
            'applicationType': 'pc',
            'operationType': 'loan',
            'applicantName': '山田 太郎',
            'department': '営業部',
            'approvedDate': '2026-09-03',
            'approvedConfirmed': True,
            'details': {
                'device_name': 'ThinkPad X1 Carbon',
                'management_number': 'PC-01234',
                'user_name': '山田 太郎',
                'operation_date': '2026-09-10',
                'expected_return_date': '2026-09-20',
                'quantity': 1,
                'location': '東京本社',
                'purpose': '営業活動',
            },
            'notes': '承認書確認済み',
        }

    def post_with_csrf(self, payload=None):
        self.client.get(reverse('applications:session'))
        token = self.client.cookies['csrftoken'].value
        return self.client.post(
            self.url,
            data=json.dumps(payload or self.payload),
            content_type='application/json',
            HTTP_X_CSRFTOKEN=token,
        )

    def test_create_list_and_sync_approved_application(self):
        response = self.post_with_csrf()
        self.assertEqual(response.status_code, 201)
        self.assertEqual(ApprovedApplication.objects.count(), 1)
        application = ApprovedApplication.objects.get()
        self.assertEqual(application.entered_by, self.user)
        self.assertEqual(response.json()['referenceNumber'], application.reference_number)
        self.assertTrue(response.json()['ledgerSynced'])
        self.assertTrue((self.ledger_directory / '承認済み_PC管理台帳.xlsx').exists())

        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()), 1)

    def test_reject_unapproved_entry(self):
        self.payload['approvedConfirmed'] = False
        response = self.post_with_csrf()
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ApprovedApplication.objects.count(), 0)


class LedgerSyncTests(TestCase):
    def setUp(self):
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.ledger_directory = Path(self.temporary_directory.name)
        self.ledger_settings = override_settings(
            LEDGER_OUTPUT_DIR=self.ledger_directory,
        )
        self.ledger_settings.enable()
        self.addCleanup(self.ledger_settings.disable)

        PcLoanApplication.objects.create(
            **COMMON_FIELDS,
            applicant_name='佐藤 花子',
            management_number='PC-00001',
            start_date=date(2026, 9, 10),
            location='東京本社',
        )
        ExternalStorageLoanApplication.objects.create(
            **COMMON_FIELDS,
            applicant_name='佐藤 花子',
            device_name='USBメモリ',
            capacity='64GB',
            location='第2会議室',
            loan_date=date(2026, 9, 10),
        )
        LanEquipmentLoanApplication.objects.create(
            **COMMON_FIELDS,
            device_type=LanEquipmentLoanApplication.DeviceType.LAN_CABLE,
            device_name='CAT6 LANケーブル',
            start_date=date(2026, 9, 10),
            return_date=date(2026, 9, 12),
            location='第2会議室',
        )
        SmartphonePurchaseApplication.objects.create(
            **COMMON_FIELDS,
            model_name='iPhone 16',
            purchase_date=date(2026, 9, 10),
            storage=SmartphonePurchaseApplication.Storage.GB_128,
            sim_required=SmartphonePurchaseApplication.SimRequired.YES,
        )

    def test_sync_all_ledgers_outputs_only_form_fields(self):
        sync_all_ledgers()

        expected_headers = {
            'PC貸出管理台帳.xlsx': (
                '申請者氏名', '所属部署', '社員番号', '貸出者氏名',
                '管理番号', '利用開始日', '利用場所',
            ),
            '外部記憶装置貸出管理台帳.xlsx': (
                '申請者氏名', '所属部署', '社員番号', '貸出者氏名',
                '機器名', '容量', '場所', '貸し出し日',
            ),
            'LAN機器貸出管理台帳.xlsx': (
                '申請者氏名', '所属部署', '社員番号', '機器種別',
                '機器名', '利用開始日', '返却予定日', '利用場所',
            ),
            'スマートフォン購入管理台帳.xlsx': (
                '申請者氏名', '所属部署', '社員番号', '機種',
                '購入日', '容量', 'SIMの有無',
            ),
        }

        for file_name, expected in expected_headers.items():
            with self.subTest(file_name=file_name):
                workbook = load_workbook(self.ledger_directory / file_name)
                worksheet = workbook.active
                headers = tuple(cell.value for cell in worksheet[1])
                self.assertEqual(headers, expected)
                self.assertEqual(worksheet.max_row, 2)
                self.assertNotIn('ID', headers)
                self.assertNotIn('申請状態', headers)
                self.assertNotIn('作成日時', headers)
                self.assertNotIn('更新日時', headers)
