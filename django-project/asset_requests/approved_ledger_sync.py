import json
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from threading import Lock

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import ApprovedApplication


DETAIL_COLUMNS = {
    "pc": (
        ("利用者氏名", "user_name"),
        ("管理番号", "management_number"),
        ("利用開始日", "start_date"),
        ("利用場所", "location"),
        ("利用目的", "purpose"),
    ),
    "memory": (
        ("利用者氏名", "user_name"),
        ("機器名", "device_name"),
        ("容量", "capacity"),
        ("貸出日", "loan_date"),
        ("利用場所", "location"),
        ("利用目的", "purpose"),
    ),
    "lan": (
        ("機器種別", "device_type"),
        ("機器名", "device_name"),
        ("必要個数", "quantity"),
        ("利用開始日", "start_date"),
        ("返却予定日", "return_date"),
        ("利用場所", "location"),
        ("利用目的", "purpose"),
    ),
    "phone": (
        ("OS", "os"),
        ("回線区分", "line_type"),
        ("機種", "model_name"),
        ("台数", "quantity"),
        ("購入日", "purchase_date"),
        ("容量", "storage"),
        ("SIM", "sim_required"),
        ("利用目的", "purpose"),
    ),
    "other": (("転記内容", "_all"),),
}

FILE_NAMES = {
    "pc": "承認済み_PC貸出管理台帳.xlsx",
    "memory": "承認済み_外部記憶装置貸出管理台帳.xlsx",
    "lan": "承認済み_LAN機器貸出管理台帳.xlsx",
    "phone": "承認済み_スマートフォン購入管理台帳.xlsx",
    "other": "承認済み_その他申請管理台帳.xlsx",
}

_SYNC_LOCK = Lock()


def _safe_value(value):
    if isinstance(value, bool):
        return "あり" if value else "なし"
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _excel_datetime(value):
    """Excelが扱えないタイムゾーン情報を外し、設定中の現地時刻にする。"""

    if timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def sync_approved_ledger(application_type):
    columns = DETAIL_COLUMNS.get(application_type)
    if columns is None:
        raise ValueError(f"未対応の申請種別です: {application_type}")

    output_dir = Path(settings.APPROVED_LEDGER_OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / FILE_NAMES[application_type]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "転記データ"
    worksheet.freeze_panes = "A2"

    common_headers = ("受付番号", "申請者氏名", "所属部署", "承認日", "元PDF", "登録担当者", "登録日時")
    worksheet.append([*common_headers, *(label for label, _key in columns), "担当者メモ"])
    header_fill = PatternFill(fill_type="solid", fgColor="17376D")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    queryset = (
        ApprovedApplication.objects.filter(application_type=application_type)
        .select_related("entered_by")
        .order_by("pk")
    )
    for record in queryset:
        detail_values = []
        for _label, key in columns:
            value = record.details if key == "_all" else record.details.get(key, "")
            detail_values.append(_safe_value(value))
        worksheet.append([
            record.reference_number,
            _safe_value(record.applicant_name),
            _safe_value(record.department),
            record.approved_date,
            Path(record.source_pdf.name).name,
            _safe_value(record.entered_by.get_username()),
            _excel_datetime(record.created_at),
            *detail_values,
            _safe_value(record.notes),
        ])

    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        longest = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(longest + 3, 42)

    with _SYNC_LOCK:
        temporary_path = None
        try:
            with NamedTemporaryFile(
                dir=output_dir,
                prefix=f".{output_path.stem}-",
                suffix=".tmp.xlsx",
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            os.replace(temporary_path, output_path)
        finally:
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()

    return output_path
