from datetime import date, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.admin.sites import AdminSite
from django.contrib.auth import get_user_model
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient, APITestCase
from openpyxl import load_workbook

from accounts.models import Department, UserProfile

from .admin import PCRequestAdmin
from .models import (
    ApprovedApplication,
    ExternalStorageRequest,
    LANRequest,
    PCRequest,
    RequestStatus,
    SmartphoneRequest,
)


def relative_date(days_from_today):
    return (timezone.localdate() + timedelta(days=days_from_today)).isoformat()


def expected_reference_number(asset_request):
    created_date = timezone.localdate(asset_request.created_at)
    return (
        f"{asset_request.REFERENCE_PREFIX}-{created_date:%Y%m%d}-"
        f"{asset_request.pk:06d}"
    )


class AuthenticatedAssetRequestAPITestCase(APITestCase):
    def setUp(self):
        super().setUp()
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.ledger_directory = Path(self.temporary_directory.name)
        self.ledger_settings = override_settings(
            LEDGER_OUTPUT_DIR=self.ledger_directory,
        )
        self.ledger_settings.enable()
        self.addCleanup(self.ledger_settings.disable)
        self.user = get_user_model().objects.create_user(
            username="asset-request-user",
            email="asset-request@example.com",
            password="Test-password-123!",
        )
        self.profile = UserProfile.objects.create(
            user=self.user,
            display_name="プロフィール 太郎",
            department=Department.SYSTEM,
        )
        self.client.force_login(self.user)


class ReferenceNumberAdminTests(APITestCase):
    def test_request_can_be_found_by_its_complete_reference_number(self):
        user = get_user_model().objects.create_user(
            username="admin-search-user",
            password="Test-password-123!",
        )
        pc_request = PCRequest.objects.create(
            created_by=user,
            requester_name="管理画面テスト",
            department="情報システム部",
            requester_email="admin-search@example.com",
            applicant_name="管理画面テスト",
            management_number="PC-ADMIN-001",
            start_date=timezone.localdate(),
            location="東京本社",
            purpose="受付番号検索テスト",
        )
        model_admin = PCRequestAdmin(PCRequest, AdminSite())

        queryset, use_distinct = model_admin.get_search_results(
            request=None,
            queryset=PCRequest.objects.all(),
            search_term=pc_request.reference_number.lower(),
        )

        self.assertQuerySetEqual(queryset, [pc_request])
        self.assertFalse(use_distinct)

    def test_status_can_be_edited_from_admin_list(self):
        model_admin = PCRequestAdmin(PCRequest, AdminSite())

        self.assertIn("status", model_admin.list_editable)


class RequestStatusModelTests(APITestCase):
    def test_status_update_is_persisted(self):
        user = get_user_model().objects.create_user(
            username="status-user",
            password="Test-password-123!",
        )
        pc_request = PCRequest.objects.create(
            created_by=user,
            requester_name="状態更新テスト",
            department="情報システム部",
            requester_email="status@example.com",
            applicant_name="状態更新テスト",
            management_number="PC-STATUS-001",
            start_date=timezone.localdate(),
            location="東京本社",
            purpose="申請状態の更新確認",
        )

        self.assertEqual(pc_request.status, RequestStatus.PENDING)

        pc_request.status = RequestStatus.APPROVED
        pc_request.save(update_fields=["status"])
        pc_request.refresh_from_db()

        self.assertEqual(pc_request.status, RequestStatus.APPROVED)


class PCRequestCreateAPITests(AuthenticatedAssetRequestAPITestCase):
    url_name = "asset_requests:pc-request-create"

    def get_payload(self):
        return {
            "applicant_name": "テスト太郎",
            "management_number": "PC-TEST-001",
            "start_date": relative_date(1),
            "location": "テスト環境",
            "purpose": "API自動テスト",
        }

    def test_pc_request_can_be_created(self):
        response = self.client.post(
            reverse(self.url_name),
            self.get_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(PCRequest.objects.count(), 1)

        pc_request = PCRequest.objects.get()
        self.assertEqual(pc_request.created_by, self.user)
        self.assertEqual(pc_request.requester_name, self.profile.display_name)
        self.assertEqual(pc_request.department, "システム部")
        self.assertEqual(pc_request.requester_email, self.user.email)
        self.assertEqual(pc_request.purpose, "API自動テスト")
        self.assertEqual(pc_request.status, "pending")
        self.assertEqual(
            response.data["reference_number"],
            expected_reference_number(pc_request),
        )
        self.assertIs(response.data["ledger_synced"], True)
        self.assertIsNone(response.data["ledger_warning"])

        workbook = load_workbook(self.ledger_directory / "PC貸出管理台帳.xlsx")
        worksheet = workbook["PC貸出"]
        self.assertEqual(
            tuple(cell.value for cell in worksheet[1]),
            (
                "申請者氏名",
                "所属部署",
                "メールアドレス",
                "利用者氏名",
                "管理番号",
                "利用開始日",
                "利用場所",
                "利用目的",
            ),
        )
        self.assertEqual(worksheet.max_row, 2)

    def test_request_remains_saved_when_ledger_sync_fails(self):
        blocked_output_path = self.ledger_directory / "not-a-directory"
        blocked_output_path.write_text(
            "block directory creation",
            encoding="utf-8",
        )

        with self.assertLogs("asset_requests.views", level="ERROR"):
            with override_settings(LEDGER_OUTPUT_DIR=blocked_output_path):
                response = self.client.post(
                    reverse(self.url_name),
                    self.get_payload(),
                    format="json",
                )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(PCRequest.objects.count(), 1)
        self.assertIs(response.data["ledger_synced"], False)
        self.assertIsNotNone(response.data["ledger_warning"])

    def test_pc_request_requires_purpose(self):
        payload = self.get_payload()
        payload.pop("purpose")

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("purpose", response.data)
        self.assertFalse(PCRequest.objects.exists())

    def test_pc_request_status_cannot_be_set_by_client(self):
        payload = self.get_payload()
        payload["status"] = "approved"
        payload["reference_number"] = "FORGED-REFERENCE"

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(PCRequest.objects.get().status, "pending")
        self.assertEqual(response.data["status"], "pending")
        self.assertNotEqual(response.data["reference_number"], "FORGED-REFERENCE")

    def test_pc_request_creator_cannot_be_spoofed_by_client(self):
        other_user = get_user_model().objects.create_user(
            username="other-user",
            password="Test-password-123!",
        )
        payload = self.get_payload()
        payload["created_by"] = other_user.pk

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(PCRequest.objects.get().created_by, self.user)

    def test_requester_snapshot_cannot_be_spoofed_by_client(self):
        payload = self.get_payload()
        payload.update({
            "requester_name": "別人",
            "department": "不正な部署",
            "requester_email": "attacker@example.com",
        })

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        pc_request = PCRequest.objects.get()
        self.assertEqual(pc_request.requester_name, self.profile.display_name)
        self.assertEqual(pc_request.department, "システム部")
        self.assertEqual(pc_request.requester_email, self.user.email)

    def test_pc_request_rejects_past_start_date(self):
        payload = self.get_payload()
        payload["start_date"] = relative_date(-1)

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("start_date", response.data)
        self.assertFalse(PCRequest.objects.exists())


class ExternalStorageRequestCreateAPITests(AuthenticatedAssetRequestAPITestCase):
    url_name = "asset_requests:external-storage-request-create"

    def get_payload(self):
        return {
            "applicant_name": "テスト花子",
            "device_name": "暗号化USBメモリ",
            "capacity": "64GB",
            "location": "東京本社",
            "loan_date": relative_date(2),
            "purpose": "データ受け渡し",
        }

    def test_external_storage_request_can_be_created(self):
        response = self.client.post(
            reverse(self.url_name),
            self.get_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(ExternalStorageRequest.objects.count(), 1)

        request = ExternalStorageRequest.objects.get()
        self.assertEqual(request.created_by, self.user)
        self.assertEqual(request.device_name, "暗号化USBメモリ")
        self.assertEqual(request.purpose, "データ受け渡し")
        self.assertEqual(request.status, "pending")
        self.assertEqual(
            response.data["reference_number"],
            expected_reference_number(request),
        )

    def test_external_storage_request_requires_purpose(self):
        payload = self.get_payload()
        payload.pop("purpose")

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("purpose", response.data)
        self.assertFalse(ExternalStorageRequest.objects.exists())

    def test_external_storage_request_rejects_past_loan_date(self):
        payload = self.get_payload()
        payload["loan_date"] = relative_date(-1)

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("loan_date", response.data)
        self.assertFalse(ExternalStorageRequest.objects.exists())


class LANRequestCreateAPITests(AuthenticatedAssetRequestAPITestCase):
    url_name = "asset_requests:lan-request-create"

    def get_payload(self):
        return {
            "device_type": LANRequest.DeviceType.USB_LAN_ADAPTER,
            "device_name": "USB-C LANアダプター",
            "quantity": 2,
            "start_date": relative_date(3),
            "return_date": relative_date(10),
            "location": "大阪支社",
            "purpose": "出張先での有線接続",
            "notes": "USB-C対応",
        }

    def test_lan_request_can_be_created(self):
        payload = self.get_payload()
        response = self.client.post(
            reverse(self.url_name),
            payload,
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(LANRequest.objects.count(), 1)

        request = LANRequest.objects.get()
        self.assertEqual(request.created_by, self.user)
        self.assertEqual(request.quantity, 2)
        self.assertEqual(
            request.return_date,
            date.fromisoformat(payload["return_date"]),
        )
        self.assertEqual(request.status, "pending")
        self.assertEqual(
            response.data["reference_number"],
            expected_reference_number(request),
        )

    def test_lan_request_requires_purpose(self):
        payload = self.get_payload()
        payload.pop("purpose")

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("purpose", response.data)
        self.assertFalse(LANRequest.objects.exists())

    def test_lan_request_rejects_quantity_less_than_one(self):
        payload = self.get_payload()
        payload["quantity"] = 0

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("quantity", response.data)
        self.assertFalse(LANRequest.objects.exists())

    def test_lan_request_rejects_return_date_before_start_date(self):
        payload = self.get_payload()
        payload["return_date"] = relative_date(2)

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("return_date", response.data)
        self.assertFalse(LANRequest.objects.exists())

    def test_lan_request_rejects_past_start_date(self):
        payload = self.get_payload()
        payload["start_date"] = relative_date(-1)

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("start_date", response.data)
        self.assertFalse(LANRequest.objects.exists())


class SmartphoneRequestCreateAPITests(AuthenticatedAssetRequestAPITestCase):
    url_name = "asset_requests:smartphone-request-create"

    def get_payload(self):
        return {
            "os": SmartphoneRequest.OS.IOS,
            "line_type": SmartphoneRequest.LineType.NEW_CONTRACT,
            "model_name": "iPhoneテストモデル",
            "quantity": 1,
            "purchase_date": relative_date(4),
            "storage": SmartphoneRequest.Storage.GB_128,
            "sim_required": False,
            "purpose": "営業連絡",
            "notes": "テスト申請",
        }

    def test_smartphone_request_can_be_created(self):
        payload = self.get_payload()
        response = self.client.post(
            reverse(self.url_name),
            payload,
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(SmartphoneRequest.objects.count(), 1)

        request = SmartphoneRequest.objects.get()
        self.assertEqual(request.created_by, self.user)
        self.assertEqual(request.model_name, "iPhoneテストモデル")
        self.assertEqual(
            request.purchase_date,
            date.fromisoformat(payload["purchase_date"]),
        )
        self.assertIs(request.sim_required, False)
        self.assertEqual(request.status, "pending")
        self.assertEqual(
            response.data["reference_number"],
            expected_reference_number(request),
        )

    def test_smartphone_request_requires_purpose(self):
        payload = self.get_payload()
        payload.pop("purpose")

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("purpose", response.data)
        self.assertFalse(SmartphoneRequest.objects.exists())

    def test_smartphone_request_requires_sim_required(self):
        payload = self.get_payload()
        payload.pop("sim_required")

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("sim_required", response.data)
        self.assertFalse(SmartphoneRequest.objects.exists())

    def test_smartphone_request_requires_model_name_and_purchase_date(self):
        for field_name in ("model_name", "purchase_date"):
            with self.subTest(field_name=field_name):
                payload = self.get_payload()
                payload.pop(field_name)

                response = self.client.post(
                    reverse(self.url_name),
                    payload,
                    format="json",
                )

                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertIn(field_name, response.data)

        self.assertFalse(SmartphoneRequest.objects.exists())

    def test_smartphone_request_rejects_quantity_less_than_one(self):
        payload = self.get_payload()
        payload["quantity"] = 0

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("quantity", response.data)
        self.assertFalse(SmartphoneRequest.objects.exists())

    def test_smartphone_request_rejects_past_purchase_date(self):
        payload = self.get_payload()
        payload["purchase_date"] = relative_date(-1)

        response = self.client.post(reverse(self.url_name), payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("purchase_date", response.data)
        self.assertFalse(SmartphoneRequest.objects.exists())


class CreateEndpointMethodTests(AuthenticatedAssetRequestAPITestCase):
    def test_create_endpoints_do_not_expose_request_lists(self):
        url_names = [
            "asset_requests:pc-request-create",
            "asset_requests:external-storage-request-create",
            "asset_requests:lan-request-create",
            "asset_requests:smartphone-request-create",
        ]

        for url_name in url_names:
            with self.subTest(url_name=url_name):
                response = self.client.get(reverse(url_name))

                self.assertEqual(
                    response.status_code,
                    status.HTTP_405_METHOD_NOT_ALLOWED,
                )


class AssetRequestAuthenticationTests(APITestCase):
    def test_anonymous_users_cannot_create_asset_requests(self):
        url_names = [
            "asset_requests:pc-request-create",
            "asset_requests:external-storage-request-create",
            "asset_requests:lan-request-create",
            "asset_requests:smartphone-request-create",
        ]

        for url_name in url_names:
            with self.subTest(url_name=url_name):
                response = self.client.post(reverse(url_name), {}, format="json")

                self.assertEqual(
                    response.status_code,
                    status.HTTP_403_FORBIDDEN,
                )

        self.assertFalse(PCRequest.objects.exists())
        self.assertFalse(ExternalStorageRequest.objects.exists())
        self.assertFalse(LANRequest.objects.exists())
        self.assertFalse(SmartphoneRequest.objects.exists())


class AssetRequestProfileTests(APITestCase):
    def test_authenticated_user_must_complete_profile_before_requesting(self):
        user = get_user_model().objects.create_user(
            username="no-profile-user",
            email="no-profile@example.com",
            password="Test-password-123!",
        )
        self.client.force_login(user)

        response = self.client.post(
            reverse("asset_requests:pc-request-create"),
            {
                "applicant_name": "プロフィール未登録",
                "management_number": "PC-NO-PROFILE",
                "start_date": relative_date(1),
                "location": "東京本社",
                "purpose": "プロフィール必須テスト",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("profile", response.data)
        self.assertFalse(PCRequest.objects.exists())


class AssetRequestCSRFProtectionTests(APITestCase):
    url_name = "asset_requests:pc-request-create"
    password = "Test-password-123!"

    def setUp(self):
        super().setUp()
        self.user = get_user_model().objects.create_user(
            username="asset-csrf-user",
            email="asset-csrf@example.com",
            password=self.password,
        )
        UserProfile.objects.create(
            user=self.user,
            display_name="CSRFテスト太郎",
            department=Department.SALES,
        )
        self.csrf_client = APIClient(enforce_csrf_checks=True)
        self.csrf_client.force_login(self.user)

    def get_payload(self):
        return {
            "applicant_name": "CSRFテスト太郎",
            "management_number": "PC-CSRF-001",
            "start_date": relative_date(1),
            "location": "テスト環境",
            "purpose": "CSRF自動テスト",
        }

    def test_authenticated_request_without_csrf_token_is_rejected(self):
        response = self.csrf_client.post(
            reverse(self.url_name),
            self.get_payload(),
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertFalse(PCRequest.objects.exists())

    def test_authenticated_request_with_valid_csrf_token_is_accepted(self):
        self.csrf_client.get(reverse("accounts:session"))
        csrf_token = self.csrf_client.cookies["csrftoken"].value

        response = self.csrf_client.post(
            reverse(self.url_name),
            self.get_payload(),
            format="json",
            HTTP_X_CSRFTOKEN=csrf_token,
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(PCRequest.objects.get().created_by, self.user)


class ApprovedApplicationAPITests(APITestCase):
    def setUp(self):
        super().setUp()
        self.temporary_directory = TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        temporary_path = Path(self.temporary_directory.name)
        self.settings_override = override_settings(
            APPROVED_LEDGER_OUTPUT_DIR=temporary_path / "ledgers",
        )
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.ledger_directory = temporary_path / "ledgers"
        self.user = get_user_model().objects.create_user(
            username="entry-operator",
            password="Test-password-123!",
        )
        self.client.force_login(self.user)
        self.url = reverse("asset_requests:approved-application-list-create")

    def get_payload(self):
        return {
            "application_type": "pc",
            "operation_type": "loan",
            "applicant_name": "申請 太郎",
            "department": "営業部",
            "approved_date": timezone.localdate().isoformat(),
            "details": {
                "device_name": "ノートPC",
                "user_name": "利用 花子",
                "management_number": "PC-001",
                "operation_date": timezone.localdate().isoformat(),
                "expected_return_date": relative_date(30),
                "quantity": 1,
                "location": "東京本社",
                "purpose": "顧客訪問",
            },
            "notes": "押印確認済み",
        }

    def test_operator_can_register_approved_entry_and_generate_ledger(self):
        response = self.client.post(self.url, self.get_payload(), format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        record = ApprovedApplication.objects.get()
        self.assertEqual(record.entered_by, self.user)
        self.assertIs(response.data["ledger_synced"], True)
        self.assertTrue(
            (self.ledger_directory / "承認済み_PC管理台帳.xlsx").exists()
        )

    def test_purchase_does_not_require_management_number(self):
        payload = self.get_payload()
        payload["operation_type"] = "purchase"
        payload["details"] = {
            "device_name": "ノートPC",
            "operation_date": timezone.localdate().isoformat(),
            "quantity": 2,
            "purpose": "新入社員用",
        }

        response = self.client.post(self.url, payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertNotIn("management_number", ApprovedApplication.objects.get().details)

    def test_loan_requires_management_number(self):
        payload = self.get_payload()
        details = payload["details"]
        details.pop("management_number")

        response = self.client.post(self.url, payload, format="json")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(ApprovedApplication.objects.exists())

    def test_anonymous_user_cannot_access_records(self):
        self.client.logout()

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
